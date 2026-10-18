"""
Shared bits for the Excel export: colours, labels and small helper functions.
Used by excel_export.py (the Excel file the app really builds) and excel_dashboard.py (an alternative dashboard that is not connected yet).
"""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Colours match the web app (plum + peach on cream).
PLUM, PEACH, CREAM, INK = "5C3A41", "ECA087", "F5F1EA", "2B1B17"
MUTED = "6B5A50"
RED_FILL, RED_TEXT = "F6D5D5", "A32A2A"
GREEN_TEXT = "2F6B3A"

FLAG = "\U0001F6A9"   # red flag emoji, put after the name of a suspicious participant
BANDS = ["Excellent", "Very good", "Good", "Need improvement", "Pending review"]
DIMS = [("semantic", "Semantic"), ("stylistic", "Stylistic"), ("pragmatic", "Pragmatic")]

_thin = Side(style="thin", color="D9D0C3")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
HEAD_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
HEAD_FILL = PatternFill("solid", fgColor=PLUM)


def num(value):
    """A number, or None when the value is empty / not a number."""
    try:
        if value is None or value == "":
            return None
        return float(value)
    except (TypeError, ValueError):
        return None


def mean(values):
    """The average of the numbers in `values` (ignoring empty ones), or None if there are none."""
    nums = [n for n in (num(v) for v in values) if n is not None]
    return round(sum(nums) / len(nums), 2) if nums else None


def set_text(cell, value):
    """Write text that must never be treated as a formula (names like '=SUM(...)')."""
    cell.value = str(value)
    cell.data_type = "s"


def head_cell(ws, row, col, label):
    """One purple header cell with white bold text."""
    c = ws.cell(row=row, column=col, value=label)
    c.font, c.fill, c.border = HEAD_FONT, HEAD_FILL, BORDER
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def header_row(ws, row, labels):
    """A whole header row, starting in column A."""
    for col, label in enumerate(labels, start=1):
        head_cell(ws, row, col, label)
    ws.row_dimensions[row].height = 30


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w